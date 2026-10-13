"""Small synthetic Office fixtures, no model or network requests."""

import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document as WordDocument
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication

from read_document import load_imported_text, collect_import_paths, unique_source_name
from split_documents import split_documents
from desktop_app import MainWindow


class OfficeImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        replacement = patch("desktop_app.default_material_path", return_value=self.folder / "materials.json")
        replacement.start()
        self.addCleanup(replacement.stop)

    def make_word(self, name="test.docx", empty=False):
        document = WordDocument()
        if not empty:
            document.add_paragraph("项目代号是青竹。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "负责人"
            table.cell(0, 1).text = "林晓"
            table.cell(1, 0).text = "预算"
            table.cell(1, 1).text = "300"
            document.add_paragraph("项目位于武汉。")
        path = self.folder / name
        document.save(path)
        return path

    def make_excel(self, empty=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "项目资料"
        if not empty:
            sheet.append(["代号", "青竹", 0, False])
            sheet.append(["预算", 300, "=B2*2"])
            workbook.create_sheet("负责人").append(["姓名", "林晓"])
        path = self.folder / "test.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def test_word_body_table_order_and_source(self):
        document = load_imported_text(self.make_word("中文.DOCX"))
        text = document.page_content
        self.assertLess(text.index("青竹"), text.index("林晓"))
        self.assertLess(text.index("林晓"), text.index("武汉"))
        self.assertIn("负责人 | 林晓", text)
        self.assertEqual(document.metadata["source"], "中文.DOCX")
        chunks = split_documents([document])
        for chunk in chunks:
            start = chunk.metadata["start_index"]
            self.assertEqual(text[start:start + len(chunk.page_content)], chunk.page_content)
            self.assertEqual(chunk.metadata["source"], "中文.DOCX")

    def test_excel_sheets_coordinates_and_formula_not_result(self):
        document = load_imported_text(self.make_excel())
        text = document.page_content
        for expected in ["工作表: 项目资料", "第 1 行", "B1=青竹", "C1=0", "D1=False",
                         "公式（未计算）: =B2*2", "工作表: 负责人", "B1=林晓"]:
            self.assertIn(expected, text)
        self.assertNotIn("600", text)

    def test_empty_office_rejected(self):
        for path in [self.make_word(empty=True), self.make_excel(empty=True)]:
            with self.subTest(path=path), self.assertRaisesRegex(ValueError, "正文为空"):
                load_imported_text(path)

    def test_corrupt_legacy_and_wrong_format_rejected(self):
        for name in ["broken.docx", "broken.xlsx", "old.doc", "old.xls"]:
            path = self.folder / name
            path.write_bytes(b"not an office file")
            with self.subTest(name=name), self.assertRaises(ValueError):
                load_imported_text(path)
        wrong = self.make_word().with_suffix(".xlsx")
        wrong.write_bytes(self.make_word().read_bytes())
        with self.assertRaises(ValueError):
            load_imported_text(wrong)

    def test_compressed_expansion_limit(self):
        path = self.folder / "large.docx"
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", b"a" * (21 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "解压后过大"):
            load_imported_text(path)

    def test_excel_huge_sheet_rejected(self):
        workbook = Workbook()
        workbook.active["A20001"] = "out of range"
        path = self.folder / "huge.xlsx"
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "范围过大"):
            load_imported_text(path)

    def test_mixed_desktop_import_preserves_good_files(self):
        app = QApplication.instance() or QApplication([])
        window = MainWindow()
        self.addCleanup(window.close)
        txt = self.folder / "note.txt"
        txt.write_text("自编测试资料", encoding="utf-8")
        bad = self.folder / "broken.xlsx"
        bad.write_bytes(b"broken")
        window.import_files([str(path) for path in [self.make_word(), self.make_excel(), txt, bad]])
        self.assertEqual(len(window.imported_documents), 3)
        self.assertTrue(window.materials_dirty)
        self.assertFalse(window.query_button.isEnabled())
        self.assertIn("broken.xlsx", window.import_errors.toPlainText())
        self.assertIn("青竹", window.imported_documents[1].page_content)

    def test_folder_scan_bounds_and_skips_linked_directories(self):
        linked = self.folder / "linked"
        linked.mkdir()
        (linked / "outside.txt").write_text("不读取", encoding="utf-8")
        for name in ["a.txt", "b.md"]:
            (self.folder / name).write_text("资料", encoding="utf-8")
        with patch.object(Path, "is_junction", autospec=True, side_effect=lambda p: p == linked):
            self.assertEqual([p.name for p in collect_import_paths(self.folder)], ["a.txt", "b.md"])
        with patch("read_document.MAX_FOLDER_ENTRIES", 1), self.assertRaisesRegex(ValueError, "条目超过"):
            collect_import_paths(self.folder)
        with patch("read_document.MAX_MATERIAL_FILES", 1), self.assertRaisesRegex(ValueError, "文件超过"):
            collect_import_paths(self.folder)

    def test_source_disambiguation_preserves_suffix_and_length(self):
        path = Path("b") / "same.md"
        name = unique_source_name(path, {"same.md", "same (b).md"})
        self.assertEqual(name, "same (b, 2).md")
        long_path = Path("b" * 240) / ("a" * 240 + ".md")
        first = unique_source_name(long_path, {long_path.name.casefold()})
        second = unique_source_name(long_path, {long_path.name.casefold(), first.casefold()})
        self.assertNotEqual(first, second)
        self.assertLessEqual(len(second), 255)
        self.assertEqual(Path(second).suffix, ".md")


if __name__ == "__main__":
    unittest.main()
