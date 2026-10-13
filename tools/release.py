"""Explicit-source release assembly and isolated executable smoke verification."""

import argparse
import hashlib
import json
import shutil
import sys
import tempfile
import traceback
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parents[1]
VERSION = "0.1.0-rc1"
MODULES = [
    "app_paths.py", "desktop_app.py", "document_tools.py", "embed_query.py",
    "index_cache.py", "local_answer.py", "material_store.py", "read_document.py",
    "split_documents.py", "compare_retrieval.py",
]


def smoke(output):
    """No API requests, no real user store, no external documents."""
    report = {"version": VERSION, "frozen": bool(getattr(sys, "frozen", False))}
    try:
        from PySide6.QtWidgets import QApplication
        from langchain_core.documents import Document
        from desktop_app import MainWindow
        from docx import Document as WordDocument
        from openpyxl import Workbook
        from app_paths import application_root
        from read_document import load_imported_text
        from embed_query import load_index, retrieve
        from local_answer import LocalAnswerSession

        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory(prefix="ai-release-") as tmp:
            tmp = Path(tmp)
            word = WordDocument()
            word.add_paragraph("项目代号是青竹。")
            word.save(tmp / "project.docx")
            sheet = Workbook()
            sheet.active.append(["项目预算", 300])
            sheet.save(tmp / "budget.xlsx")
            docs = [load_imported_text(tmp / name) for name in ("project.docx", "budget.xlsx")]
            report["office"] = [d.page_content for d in docs]
            window = MainWindow(storage_path=tmp / "materials.json")
            window.import_files([tmp / "project.docx", tmp / "budget.xlsx"])
            assert len(window.imported_documents) == 2
            chunk = Document(page_content=r"范数不超过$\theta$，方向$\mathbf{g}$。",
                             metadata={"source": "math.md", "start_index": 0})
            window.on_results(None, [(0, chunk, .5)])
            assert window.output.toPlainText() == chunk.page_content
            assert window.preview.toPlainText().count("\ufffc") == 2
            report["formula"] = True
            window.show()
            app.processEvents()
            window.grab().save(str(output.with_suffix(".png")))
            window.close()
            index = load_index("bge-m3", docs, cache_path=tmp / "index.npz", device="cpu")
            hits = retrieve("项目代号是什么？", *index, k=2)
            report["retrieval"] = [c.metadata["source"] for _, c, _ in hits]
            assert hits[0][1].metadata["source"] == "project.docx"
            report["root"] = str(application_root())
            if (application_root() / "runtime/llama-bin/llama-server.exe").is_file():
                session = LocalAnswerSession()
                try:
                    report["local"] = session.answer("项目代号是什么？", [hits[0][1]], device="cpu")
                    assert "青竹" in report["local"]["answer"]
                    assert not report["local"]["invalid_citations"]
                finally:
                    session.stop()
            else:
                raise RuntimeError("CPU llama-server is missing from release.")
        report["ok"] = True
    except Exception:
        report["ok"] = False
        report["error"] = traceback.format_exc()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1


def assemble():
    target = ROOT / "release" / VERSION
    bundle = target / "AIResearchAssistant"
    if not (bundle / "AIResearchAssistant.exe").is_file():
        raise RuntimeError("Build executable with PyInstaller first.")
    # Only selected public assets, never private materials or API configuration.
    for relative in ("models/bge-m3", "models/qwen3-4b-gguf",
                     "runtime/llama-bin", "runtime/llama-cuda13"):
        src = ROOT / relative
        if not src.is_dir():
            raise RuntimeError(f"Missing release asset: {relative}")
        shutil.copytree(src, bundle / relative, dirs_exist_ok=True)
    shutil.copy2(ROOT / "docs/RELEASE.md", bundle / "START_HERE.md")
    license_dir = bundle / "THIRD_PARTY_NOTICES"
    license_dir.mkdir(exist_ok=True)
    shutil.copy2(ROOT / "runtime/llama.cpp/LICENSE", license_dir / "llama.cpp-LICENSE")
    source_files = [ROOT / p for p in MODULES]
    source_files += [ROOT / p for p in ("README.md", "requirements.txt",
                                       "requirements-desktop.txt", "requirements-office.txt",
                                       "requirements-build.txt", ".gitignore")]
    for folder in ("tests", "docs", "tools", "packaging", "corpus", "eval"):
        source_files.extend(p for p in (ROOT / folder).rglob("*")
                            if p.is_file() and not {"__pycache__", "stress_results"}.intersection(p.parts))
    # The desktop suite uses this synthetic fixture; no other data is shipped.
    source_files += [ROOT / "data/import_qa_test.txt"]
    with ZipFile(target / f"AIResearchAssistant-{VERSION}-source.zip", "w", ZIP_DEFLATED) as archive:
        for path in sorted(set(source_files)):
            archive.write(path, path.relative_to(ROOT))
    manifests = []
    for path in sorted(bundle.rglob("*")):
        if path.is_file():
            digest = hashlib.sha256()
            with path.open("rb") as stream:
                for block in iter(lambda: stream.read(4 * 1024 * 1024), b""):
                    digest.update(block)
            manifests.append({"path": path.relative_to(bundle).as_posix(),
                              "bytes": path.stat().st_size, "sha256": digest.hexdigest()})
    (target / "manifest.json").write_text(json.dumps(manifests, indent=2), encoding="utf-8")
    print(f"Assembled {target} ({len(manifests)} files).")

def archive_bundle():
    target = ROOT / "release" / VERSION
    bundle = target / "AIResearchAssistant"
    report = json.loads((target / "executable-smoke.json").read_text(encoding="utf-8"))
    if not report.get("ok") or not report.get("frozen"):
        raise RuntimeError("Executable smoke test must pass before archiving.")
    destination = target / f"AIResearchAssistant-{VERSION}-windows-x64.zip"
    with ZipFile(destination, "w", ZIP_DEFLATED, compresslevel=1) as archive:
        for path in sorted(bundle.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(target))
    digest = hashlib.sha256()
    with destination.open("rb") as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(block)
    destination.with_suffix(".zip.sha256").write_text(
        digest.hexdigest() + "  " + destination.name + "\n", encoding="ascii")
    with ZipFile(destination) as archive:
        forbidden = [n for n in archive.namelist()
                     if Path(n).name in {".env", "materials.json", "index.npz"}
                     or "/data/stress/" in n]
        if forbidden:
            raise RuntimeError(f"Forbidden release files: {forbidden}")
    print(f"Archived {destination} ({destination.stat().st_size} bytes).")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--assemble", action="store_true")
    parser.add_argument("--archive", action="store_true")
    args = parser.parse_args()
    if args.assemble:
        assemble()
    if args.archive:
        archive_bundle()
